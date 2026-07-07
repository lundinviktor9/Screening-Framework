#!/usr/bin/env python3
"""apply_assumptions.py - deterministic three-tier underwrite-assumptions applier.

The generalised replacement for the old flag-per-unit sign-off. The analyst fills a small
structured `underwrite_assumptions` block ONCE per deal (see
underwrite/schemas/underwrite_assumptions.schema.json); this module merges it into a copy of the
DealRR sheet BEFORE inject_deal_v21 runs, so the values cascade through the existing header-driven
template wiring (no template surgery needed).

Precedence is most-specific-wins:  unit override  >  asset override  >  deal default.

Two write modes:
  * OVERRIDE  (entry_yield, exit_yield) - a deliberate pricing call; written to EVERY unit in
    scope, replacing the schedule value.
  * FILL_BLANK (NEF, ERV growth, void, rent-free, capex, rateable, rates, service charge,
    term certain) - only populates cells the broker left EMPTY; never overwrites broker data.

Per-unit pick-list calls:
  * break_taken       -> DealRR M = 1 and Event @ Expiry (O) = "X" (vacate + re-let).
  * vacate_at_expiry  -> DealRR Event @ Expiry (O) = "X".

Deal-level erv_growth ALSO returns a Global Assumptions cell to set (D49, universal rental
growth) so growth is driven centrally, matching the Newbury v21 model.

NEVER fabricates: a field left null in every applicable tier leaves the schedule cell untouched
(a blank stays blank and is surfaced by normalise_auto.summarise_missing).

Pure openpyxl; no LibreOffice. Idempotent - re-running with the same inputs is a no-op on values.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import column_index_from_string as cix

# canonical assumption field -> DealRR column letter (field-dictionary OLD-letter scheme).
FIELD_COL = {
    "nef": "Q", "erv_growth": "R", "entry_yield": "H", "exit_yield": "I",
    "rateable_value": "V", "business_rates": "W", "service_charge": "X",
    "void_mths": "Y", "rent_free_mths": "Z", "reletting_capex_psf": "AA",
    "term_certain_mths": "BC",
}
OVERRIDE_FIELDS = {"entry_yield", "exit_yield"}          # written to every unit in scope
FILL_BLANK_FIELDS = set(FIELD_COL) - OVERRIDE_FIELDS     # only fill empty cells
PCT_FIELDS = {"nef", "erv_growth", "entry_yield", "exit_yield"}
GA_WIRING = {"erv_growth": "D49"}                        # deal-level growth -> Global Assumptions

# DealRR structural columns
ASSET_COL, UNIT_COL, EVENT_COL, BREAK_COL = "B", "E", "O", "M"


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _num(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _as_decimal(field: str, v: Any) -> Any:
    """Percentage-tolerant read for rate-like fields: 6.25 -> 0.0625, 0.0625 -> 0.0625.
    Leaves a value already in decimals untouched. Non-pct fields pass through unchanged."""
    if field not in PCT_FIELDS:
        return v
    n = _num(v)
    if n is None:
        return v
    return n / 100.0 if n > 1.0 else n


def _unit_keys(asset: str, unit: str) -> List[str]:
    """Keys under which a per-unit override may be stored (most specific first)."""
    return [f"{asset}::{unit}", str(unit)]


def resolve(field: str, asset: str, unit: str,
            deal: Dict[str, Any], assets: Dict[str, Any], units: Dict[str, Any]) -> Any:
    """Most-specific-wins lookup for one field on one unit. Returns None if unset in every tier."""
    for key in _unit_keys(asset, unit):
        u = units.get(key)
        if u is not None and u.get(field) is not None:
            return u[field]
    a = assets.get(asset)
    if a is not None and a.get(field) is not None:
        return a[field]
    if deal.get(field) is not None:
        return deal[field]
    return None


def apply_assumptions(rr_xlsx: str, assumptions: Optional[Dict[str, Any]], dest: str,
                      rr_sheet: str = "DealRR") -> Dict[str, Any]:
    """Merge the three-tier assumptions into a copy of the DealRR sheet, saved to `dest`.

    Returns {"notes": [...audit strings...], "ga_cells": {"D49": <growth>, ...}, "changed": <int>}.
    `ga_cells` are Global Assumptions cells the CALLER should stamp onto the model after injection
    (adapter does this). Safe to call with assumptions=None (straight copy)."""
    wb = openpyxl.load_workbook(rr_xlsx)
    if rr_sheet not in wb.sheetnames:
        raise ValueError(f"RR sheet {rr_sheet!r} not in {rr_xlsx}")
    ws = wb[rr_sheet]

    a = assumptions or {}
    deal: Dict[str, Any] = dict(a.get("deal") or {})
    assets: Dict[str, Any] = dict(a.get("assets") or {})
    units: Dict[str, Any] = dict(a.get("units") or {})

    notes: List[str] = []
    changed = 0

    # data rows: Unit Number (col E) non-empty
    data_rows = [r for r in range(2, ws.max_row + 1)
                 if ws.cell(r, cix(UNIT_COL)).value not in (None, "")]

    # tally of writes per field for a compact audit note
    tally: Dict[str, int] = {}

    def _bump(field: str) -> None:
        tally[field] = tally.get(field, 0) + 1

    for r in data_rows:
        asset = str(ws.cell(r, cix(ASSET_COL)).value or "")
        unit = str(ws.cell(r, cix(UNIT_COL)).value or "")

        # ---- valued fields (deal/asset/unit tiers) ----
        for field, col in FIELD_COL.items():
            val = resolve(field, asset, unit, deal, assets, units)
            if val is None:
                continue
            val = _as_decimal(field, val)
            c = ws.cell(r, cix(col))
            if field in OVERRIDE_FIELDS:
                if c.value != val:
                    c.value = val
                    changed += 1
                    _bump(field)
            else:  # FILL_BLANK
                if _is_blank(c.value):
                    c.value = val
                    changed += 1
                    _bump(field)

        # ---- per-unit pick-list flags (break / vacate) ----
        u_over: Dict[str, Any] = {}
        for key in _unit_keys(asset, unit):
            if key in units:
                u_over = units[key]
                break
        if u_over.get("break_taken") is True:
            if ws.cell(r, cix(BREAK_COL)).value != 1:
                ws.cell(r, cix(BREAK_COL)).value = 1
                changed += 1
                _bump("break_taken")
            if ws.cell(r, cix(EVENT_COL)).value != "X":
                ws.cell(r, cix(EVENT_COL)).value = "X"
                changed += 1
        if u_over.get("vacate_at_expiry") is True:
            if ws.cell(r, cix(EVENT_COL)).value != "X":
                ws.cell(r, cix(EVENT_COL)).value = "X"
                changed += 1
                _bump("vacate_at_expiry")

    for field, n in sorted(tally.items()):
        notes.append(f"assumptions: applied {field} to {n} unit(s)")

    # ---- Global Assumptions cells to set centrally (deal tier) ----
    ga_cells: Dict[str, Any] = {}
    for field, cell in GA_WIRING.items():
        if deal.get(field) is not None:
            ga_cells[cell] = _as_decimal(field, deal[field])
            notes.append(f"assumptions: deal {field} -> Global Assumptions!{cell} = {ga_cells[cell]}")

    wb.save(dest)
    return {"notes": notes, "ga_cells": ga_cells, "changed": changed}


if __name__ == "__main__":  # tiny CLI for manual runs
    import argparse, json
    ap = argparse.ArgumentParser()
    ap.add_argument("rr_xlsx")
    ap.add_argument("--assumptions", required=True, help="path to a JSON assumptions block")
    ap.add_argument("--out", required=True)
    ap.add_argument("--rr-sheet", default="DealRR")
    args = ap.parse_args()
    block = json.loads(Path(args.assumptions).read_text(encoding="utf-8"))
    res = apply_assumptions(args.rr_xlsx, block, args.out, args.rr_sheet)
    print(json.dumps(res, indent=2, default=str))
